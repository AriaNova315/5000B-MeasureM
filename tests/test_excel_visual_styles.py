import asyncio
import io
import tempfile
import unittest
from pathlib import Path

import aiosqlite
from openpyxl import load_workbook

import app.database as database
from app.services.excel_builder.builder import build_excel
from app.services.importer import import_project


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_INPUT = ROOT / "项目相关信息表.xlsx"


def fill_rgb(cell):
    return cell.fill.fgColor.rgb


class ExcelVisualStyleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.old_db_path = database.DB_PATH
        database.DB_PATH = str(Path(cls.temp_dir.name) / "measurem.db")
        cls.phase_workbooks = asyncio.run(cls._build_workbooks())

    @classmethod
    async def _build_workbooks(cls):
        await database.init_db()
        async with aiosqlite.connect(database.DB_PATH) as db:
            result = await import_project(
                db,
                SAMPLE_INPUT.read_bytes(),
                source_filename=SAMPLE_INPUT.name,
            )
            project_id = result["project_id"]
            rows = await (
                await db.execute(
                    "SELECT id,wbs FROM phases WHERE project_id=? AND wbs IN ('1','2') "
                    "ORDER BY CAST(wbs AS INTEGER)",
                    (project_id,),
                )
            ).fetchall()
            return {
                wbs: load_workbook(
                    io.BytesIO(await build_excel(db, project_id, phase_id)),
                    data_only=False,
                )
                for phase_id, wbs in rows
            }

    @classmethod
    def tearDownClass(cls):
        database.DB_PATH = cls.old_db_path
        cls.temp_dir.cleanup()

    def test_narrative_reports_use_section_label_and_body_layers(self):
        phase1 = self.phase_workbooks["1"]
        progress = phase1["001进展报告"]
        self.assertEqual("00D9E1F2", fill_rgb(progress["B3"]))
        self.assertEqual("00E7E6E6", fill_rgb(progress["B4"]))
        self.assertEqual("00FFFFFF", fill_rgb(progress["C4"]))
        self.assertEqual("medium", progress["B3"].border.left.style)
        self.assertEqual("medium", progress["I21"].border.right.style)

        measure = phase1["001测量与分析报告"]
        self.assertEqual("00D9E1F2", fill_rgb(measure["A6"]))
        self.assertEqual("00E7E6E6", fill_rgb(measure["A3"]))
        self.assertEqual("00FFFFFF", fill_rgb(measure["C3"]))

        meeting = phase1["LH001项目例会单"]
        self.assertEqual("00E7E6E6", fill_rgb(meeting["A3"]))
        self.assertEqual("00FFFFFF", fill_rgb(meeting["C3"]))
        self.assertGreaterEqual(meeting.row_dimensions[6].height, 30)

    def test_grid_reports_use_hierarchical_headers_and_readable_rows(self):
        phase2 = self.phase_workbooks["2"]
        workreport = phase2["WorkReport20260112"]
        self.assertEqual("00D9E1F2", fill_rgb(workreport["A1"]))
        expected_fixed_header_merges = {
            f"{column}1:{column}3" for column in "ABCDEFGHIJK"
        }
        self.assertTrue(
            expected_fixed_header_merges.issubset(
                {str(cell_range) for cell_range in workreport.merged_cells.ranges}
            )
        )
        self.assertEqual("00D9E1F2", fill_rgb(workreport["K1"]))
        self.assertEqual("00DFE3E8", fill_rgb(workreport["L2"]))
        self.assertEqual("00FFFFFF", fill_rgb(workreport["A4"]))
        self.assertGreaterEqual(workreport.row_dimensions[4].height, 24)

        tracking = phase2["ProgressTracking-ok"]
        self.assertEqual("00D9E1F2", fill_rgb(tracking["A1"]))
        self.assertEqual("00DFE3E8", fill_rgb(tracking["A2"]))
        self.assertEqual("00D9E1F2", fill_rgb(tracking["A3"]))
        self.assertGreaterEqual(tracking.row_dimensions[4].height, 24)

    def test_later_phase_reports_receive_the_same_visual_standard(self):
        phase2 = self.phase_workbooks["2"]
        progress_name = next(
            name for name in phase2.sheetnames
            if name.startswith("002进展报告")
        )
        measure_name = next(
            name for name in phase2.sheetnames
            if name.startswith("002测量与分析报告")
        )
        self.assertEqual("00D9E1F2", fill_rgb(phase2[progress_name]["B3"]))
        self.assertEqual("00E7E6E6", fill_rgb(phase2[progress_name]["B4"]))
        self.assertEqual("00D9E1F2", fill_rgb(phase2[measure_name]["A6"]))
        self.assertEqual("00E7E6E6", fill_rgb(phase2[measure_name]["A3"]))


if __name__ == "__main__":
    unittest.main()
