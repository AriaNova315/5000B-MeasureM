import io
import tempfile
import unittest
from pathlib import Path

import aiosqlite
from openpyxl import load_workbook

import app.database as database
from app.models import NarrativeBatchItem, NarrativeBatchUpdate
from app.routers.projects import (
    api_batch_update_narratives,
    api_get_narratives,
    api_get_report_types,
)
from app.services.excel_builder.builder import build_excel
from app.services.importer import import_project


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_INPUT = ROOT / "项目相关信息表.xlsx"


class NarrativeEffectiveContentTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_db_path = database.DB_PATH
        database.DB_PATH = str(Path(self.temp_dir.name) / "measurem.db")
        await database.init_db()
        self.db = await aiosqlite.connect(database.DB_PATH)
        result = await import_project(
            self.db,
            SAMPLE_INPUT.read_bytes(),
            source_filename=SAMPLE_INPUT.name,
        )
        self.project_id = result["project_id"]

        async with self.db.execute("""
            SELECT ph.id, ph.wbs, tn.id, tn.date
            FROM phases ph JOIN tracking_nodes tn ON tn.phase_id=ph.id
            WHERE ph.project_id=? ORDER BY CAST(ph.wbs AS INTEGER), tn.date
        """, (self.project_id,)) as cur:
            rows = await cur.fetchall()
        self.phase1_id, _, self.phase1_node_id, _ = rows[0]
        phase2_rows = [row for row in rows if row[1] == "2"]
        self.phase2_id = phase2_rows[0][0]
        self.phase2_first_node_id = phase2_rows[0][2]
        self.phase2_last_node_id = phase2_rows[-1][2]
        phase4_rows = [row for row in rows if row[1] == "4"]
        self.phase4_id = phase4_rows[0][0]
        self.phase4_milestone_node_id = phase4_rows[-1][2]
        self.phase5_id = next(row[0] for row in rows if row[1] == "5")

    async def asyncTearDown(self):
        await self.db.close()
        database.DB_PATH = self.old_db_path
        self.temp_dir.cleanup()

    async def test_report_instances_match_business_meeting_names(self):
        phase1 = await api_get_report_types(self.phase1_node_id, self.db)
        phase1_meetings = [(item.id, item.label) for item in phase1 if item.id.startswith("meeting_")]
        self.assertEqual([
            ("meeting_lh001", "LH001 项目启动会"),
            ("meeting_lh002", "LH002 软件研制任务书研讨会"),
            ("meeting_lh003", "LH003 阶段例会"),
        ], phase1_meetings)

        phase2_first = await api_get_report_types(self.phase2_first_node_id, self.db)
        phase2_meetings = [(item.id, item.label) for item in phase2_first if item.id.startswith("meeting_")]
        self.assertEqual([
            ("meeting_planning", "LH004 项目策划会"),
            ("meeting_regular", "LH005 双周例会"),
        ], phase2_meetings)

        phase2_last = await api_get_report_types(self.phase2_last_node_id, self.db)
        self.assertIn(
            ("meeting_regular", "LH006 阶段例会"),
            [(item.id, item.label) for item in phase2_last],
        )

    async def test_generated_meeting_content_is_visible_without_becoming_override(self):
        sections = await api_get_narratives(
            self.phase2_first_node_id, "meeting_regular", self.db,
        )
        by_section = {section.section: section for section in sections}

        self.assertEqual("610会议室", by_section["location"].content)
        self.assertTrue(by_section["agenda"].content.startswith("1、项目组长通报"))
        self.assertIn("项目组成员报告个人工作情况", by_section["agenda"].content)
        self.assertFalse(any(section.is_overridden for section in sections))

    async def test_lh003_is_regular_phase_meeting_with_one_fewer_agenda_item(self):
        lh003 = await api_get_narratives(self.phase1_node_id, "meeting_lh003", self.db)
        agenda = next(section.content for section in lh003 if section.section == "lh003_agenda")

        self.assertEqual(5, len(agenda.splitlines()))
        self.assertNotIn("项目组成员报告个人工作情况", agenda)
        self.assertIn("质量保证人员通报", agenda)

    async def test_manual_override_controls_excel_and_empty_override_restores_default(self):
        override = NarrativeBatchUpdate(items=[
            NarrativeBatchItem(section="location", content="第一会议室"),
            NarrativeBatchItem(section="agenda", content="1、人工确认的例会议题。"),
        ])
        await api_batch_update_narratives(
            self.phase2_first_node_id, "meeting_regular", override, self.db,
        )

        edited = await api_get_narratives(
            self.phase2_first_node_id, "meeting_regular", self.db,
        )
        edited_by_section = {section.section: section for section in edited}
        self.assertTrue(edited_by_section["location"].is_overridden)
        self.assertEqual("第一会议室", edited_by_section["location"].content)

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase2_id)
        ), data_only=False)
        self.assertEqual("第一会议室", workbook["LH005项目例会单"]["G3"].value)
        self.assertEqual("1、人工确认的例会议题。", workbook["LH005项目例会单"]["C6"].value)

        reset = NarrativeBatchUpdate(items=[
            NarrativeBatchItem(section="location", content=""),
            NarrativeBatchItem(section="agenda", content=""),
        ])
        await api_batch_update_narratives(
            self.phase2_first_node_id, "meeting_regular", reset, self.db,
        )
        restored = await api_get_narratives(
            self.phase2_first_node_id, "meeting_regular", self.db,
        )
        restored_by_section = {section.section: section for section in restored}
        self.assertFalse(restored_by_section["location"].is_overridden)
        self.assertEqual("610会议室", restored_by_section["location"].content)

    async def test_planning_meeting_override_is_persisted_and_exported(self):
        await api_batch_update_narratives(
            self.phase2_first_node_id,
            "meeting_planning",
            NarrativeBatchUpdate(items=[
                NarrativeBatchItem(section="planning_host", content="策划主持人"),
            ]),
            self.db,
        )

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase2_id)
        ), data_only=False)
        self.assertEqual("策划主持人", workbook["LH004项目例会单"]["C4"].value)

    async def test_phase1_prefixed_override_controls_the_matching_meeting_only(self):
        await api_batch_update_narratives(
            self.phase1_node_id,
            "meeting_lh002",
            NarrativeBatchUpdate(items=[
                NarrativeBatchItem(section="lh002_location", content="任务书研讨室"),
            ]),
            self.db,
        )

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase1_id)
        ), data_only=False)
        self.assertEqual("610会议室", workbook["LH001项目例会单"]["G3"].value)
        self.assertEqual("任务书研讨室", workbook["LH002项目例会单"]["G3"].value)
        self.assertEqual("610会议室", workbook["LH003项目例会单"]["G3"].value)

    async def test_phase1_measure_report_uses_dedicated_layout_and_progress_formula(self):
        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase1_id)
        ), data_only=False)
        sheet = workbook["001测量与分析报告"]
        progress_formula = "='ProgressTracking-ok'!E3"

        self.assertEqual("已完成工作量", sheet["A7"].value)
        self.assertEqual(progress_formula, sheet["D7"].value)
        self.assertEqual("规模", sheet["A8"].value)
        self.assertEqual("序号", sheet["D8"].value)
        self.assertEqual("文档名称", sheet["E8"].value)
        self.assertEqual("计划规模", sheet["H8"].value)
        self.assertEqual("实际规模", sheet["I8"].value)
        self.assertEqual("规模偏差率", sheet["J8"].value)

        # 文档清单沿用其他阶段的筛选口径：当前阶段文档、评审日期为空或不晚于节点日期。
        self.assertEqual("相控阵数字样机软件研制任务书", sheet["E9"].value)
        self.assertEqual("相控阵数字样机软件外包产品技术协议", sheet["E10"].value)
        self.assertEqual(["——", "——"], [sheet["H9"].value, sheet["H10"].value])
        self.assertEqual([12, 9], [sheet["I9"].value, sheet["I10"].value])
        self.assertEqual(["——", "——"], [sheet["J9"].value, sheet["J10"].value])
        self.assertIn("A8:C10", {str(cell_range) for cell_range in sheet.merged_cells.ranges})

        self.assertEqual("需求总数", sheet["A11"].value)
        self.assertEqual("待填写", sheet["D11"].value)
        self.assertEqual("工作量", sheet["A12"].value)
        self.assertEqual("实际工时", sheet["D12"].value)
        self.assertEqual(progress_formula, sheet["G12"].value)
        self.assertEqual("过程和产品符合性", sheet["A13"].value)
        self.assertEqual("异常状态原因分析", sheet["A14"].value)
        self.assertEqual("拟制人员", sheet["A15"].value)

        values = {
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
        self.assertNotIn("测量项分析", values)
        self.assertNotIn("计划工作量", values)
        self.assertNotIn("工作量偏差", values)

        tracking = workbook["ProgressTracking-ok"]
        self.assertEqual(5, tracking.max_column)
        self.assertEqual(
            ["WBS", "资源名称", "实际开始时间", "实际完成时间"],
            [tracking.cell(1, col).value for col in range(1, 5)],
        )
        self.assertEqual("阶段跟踪（2025.12.26）", tracking["E1"].value)
        self.assertEqual("实际已完成工作量", tracking["E2"].value)
        self.assertTrue(str(tracking["E3"].value).startswith("="))
        tracking_headers = {
            cell.value for row in tracking.iter_rows(min_row=1, max_row=2)
            for cell in row if cell.value
        }
        self.assertNotIn("任务名称", tracking_headers)
        self.assertNotIn("计划开始时间", tracking_headers)
        self.assertNotIn("计划完成时间", tracking_headers)
        self.assertNotIn("计划总工作量", tracking_headers)

    async def test_phase1_progress_report_uses_dedicated_layout_and_defaults(self):
        # 兼容旧数据库：历史风险字段不再展示；节点之后评审的文档不进入叙述和规模。
        await self.db.execute("""
            INSERT OR IGNORE INTO narratives
                (tracking_node_id, report_type, section, content)
            VALUES (?, 'progress_phase1', 'risk_issues', '旧风险内容')
        """, (self.phase1_node_id,))
        await self.db.execute("""
            INSERT INTO config_items
                (project_id, phase_wbs, seq, name, item_type, actual_pages, review_date)
            VALUES (?, '1', 99, '节点后评审文档', '文档', 100, '2026-01-01')
        """, (self.project_id,))
        await self.db.commit()

        sections = await api_get_narratives(
            self.phase1_node_id, "progress_phase1", self.db,
        )
        self.assertEqual(
            ["progress_narrative", "current_summary", "next_plan", "author"],
            [section.section for section in sections],
        )
        self.assertFalse(any(section.is_overridden for section in sections))
        by_section = {section.section: section for section in sections}
        self.assertIn(
            "相控阵数字样机软件研制任务书",
            by_section["progress_narrative"].content,
        )
        self.assertIn(
            "相控阵数字样机软件外包产品技术协议",
            by_section["progress_narrative"].content,
        )
        self.assertEqual(
            "本进展期内所有任务正常推进。",
            by_section["current_summary"].content,
        )

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase1_id)
        ), data_only=False)
        sheet = workbook["001进展报告"]

        self.assertEqual(21, sheet.max_row)
        self.assertEqual(9, sheet.max_column)
        self.assertEqual("报告顺序号：", sheet["B2"].value[:6])
        self.assertEqual("right", sheet["B2"].alignment.horizontal)
        self.assertEqual("项目状态信息与度量数据", sheet["B7"].value)
        self.assertEqual("无", sheet["C8"].value)
        self.assertEqual("本进展期内任务完成情况：", sheet["C9"].value)
        self.assertEqual(by_section["progress_narrative"].content, sheet["C10"].value)
        self.assertEqual("本进展期内项目综合情况：", sheet["C11"].value)
        self.assertEqual("生成文档规模：", sheet["C12"].value)
        self.assertEqual(21, sheet["E12"].value)
        self.assertEqual("页", sheet["F12"].value)
        self.assertEqual("本进展期内所有任务正常推进。", sheet["C13"].value)
        self.assertEqual(by_section["next_plan"].content, sheet["C14"].value)
        self.assertEqual("赵要浦", sheet["C15"].value)
        self.assertTrue(sheet["B3"].fill.fgColor.rgb.endswith("D9E1F2"))
        self.assertTrue(sheet["B7"].fill.fgColor.rgb.endswith("D9E1F2"))
        self.assertAlmostEqual(16.5, sheet.column_dimensions["B"].width, places=1)
        self.assertAlmostEqual(21.66, sheet.column_dimensions["C"].width, places=1)

        merges = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
        self.assertTrue({
            "B1:I1", "B2:I2", "B3:I3", "C4:F4", "H4:I4",
            "B7:I7", "C8:I8", "B9:B12", "C10:I10",
            "C12:D12", "G12:I12", "B17:I17", "C18:H18",
        }.issubset(merges))

        self.assertEqual(
            ["测量与分析报告", "质量审核工作报告", "配置管理报告"],
            [sheet.cell(row, 3).value for row in range(19, 22)],
        )
        self.assertEqual(
            ["2025.12.26"] * 3,
            [sheet.cell(row, 9).value for row in range(19, 22)],
        )
        values = {
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
        for removed_label in [
            "风险问题说明",
            "工作量偏差超20%任务数",
            "进度偏差超20%任务数",
            "各工作产品规模偏差",
            "项目风险跟踪表",
        ]:
            self.assertNotIn(removed_label, values)

    async def test_phase1_progress_manual_override_controls_excel(self):
        override_text = "阶段1任务已按人工确认内容完成。"
        await api_batch_update_narratives(
            self.phase1_node_id,
            "progress_phase1",
            NarrativeBatchUpdate(items=[
                NarrativeBatchItem(
                    section="progress_narrative",
                    content=override_text,
                ),
            ]),
            self.db,
        )

        sections = await api_get_narratives(
            self.phase1_node_id, "progress_phase1", self.db,
        )
        progress = next(
            section for section in sections
            if section.section == "progress_narrative"
        )
        self.assertTrue(progress.is_overridden)
        self.assertEqual(override_text, progress.content)

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase1_id)
        ), data_only=False)
        self.assertEqual(override_text, workbook["001进展报告"]["C10"].value)

    async def test_non_phase1_progress_tracking_keeps_multi_node_structure(self):
        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase2_id)
        ), data_only=False)
        tracking = workbook["ProgressTracking-ok"]

        self.assertEqual(18, tracking.max_column)
        self.assertEqual("任务名称", tracking["B1"].value)
        self.assertEqual("计划总工作量", tracking["H1"].value)
        self.assertEqual("双周跟踪（2026.1.12）", tracking["I1"].value)
        self.assertEqual("阶段跟踪（2026.1.23）", tracking["M1"].value)

        progress_sheets = [
            sheet for sheet in workbook.worksheets
            if "进展报告" in sheet.title
        ]
        self.assertEqual(2, len(progress_sheets))
        for sheet in progress_sheets:
            values = {
                cell.value
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            }
            self.assertIn("风险问题说明", values)
            self.assertGreater(sheet.max_row, 21)

    async def test_report_periods_follow_node_scope(self):
        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase2_id)
        ), data_only=False)

        def report_sheet(report_name: str, variant: str):
            return next(
                sheet for sheet in workbook.worksheets
                if report_name in sheet.title and variant in sheet.title
            )

        biweekly_measure = report_sheet("测量与分析报告", "双周")
        biweekly_progress = report_sheet("进展报告", "双周")
        phase_measure = report_sheet("测量与分析报告", "（阶段）")
        phase_progress = report_sheet("进展报告", "（阶段）")

        # 双周报告保持相邻跟踪区间；当前节点是本阶段首节点，从阶段开始日计算。
        self.assertEqual("2025.12.29-2026.1.12", biweekly_measure["H5"].value)
        self.assertEqual("2025.12.29-2026.1.12", biweekly_progress["H6"].value)

        # 阶段报告覆盖整个阶段，不从上一个双周节点开始。
        self.assertEqual("2025.12.29-2026.1.23", phase_measure["H5"].value)
        self.assertEqual("2025.12.29-2026.1.23", phase_progress["H6"].value)

        # 阶段不再自动附带里程碑报告。
        self.assertFalse(any("里程碑" in sheet.title for sheet in workbook.worksheets))

        milestone_workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase4_id)
        ), data_only=False)
        milestone_measure = next(
            sheet for sheet in milestone_workbook.worksheets
            if "测量与分析报告" in sheet.title and "里程碑" in sheet.title
        )
        milestone_progress = next(
            sheet for sheet in milestone_workbook.worksheets
            if "进展报告" in sheet.title and "里程碑" in sheet.title
        )
        self.assertEqual("2025.12.15-2026.3.12", milestone_measure["H5"].value)
        self.assertEqual("2025.12.15-2026.3.12", milestone_progress["H6"].value)

    async def test_milestone_measure_report_accumulates_prior_phases(self):
        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase4_id)
        ), data_only=False)
        milestone = next(
            sheet for sheet in workbook.worksheets
            if "测量与分析报告" in sheet.title and "里程碑" in sheet.title
        )
        phase = next(
            sheet for sheet in workbook.worksheets
            if "测量与分析报告" in sheet.title and "（阶段）" in sheet.title
        )

        # 阶段报告只列本阶段；显式里程碑报告累计截止WBS4的配置项。
        self.assertEqual(11, phase["B8"].value)
        self.assertEqual(
            ["相控阵数字样机软件研制任务书", "相控阵数字样机软件外包产品技术协议"],
            [milestone.cell(row, 3).value for row in range(8, 10)],
        )
        self.assertEqual(["——", "——"], [milestone["G8"].value, milestone["G9"].value])
        self.assertEqual([12, 9], [milestone["H8"].value, milestone["H9"].value])
        self.assertEqual(["——", "——"], [milestone["I8"].value, milestone["I9"].value])

        # 阶段1的实际工作量57补入计划口径，并分别累计四类工作量/进度指标。
        document_names = [
            milestone.cell(row, 3).value
            for row in range(8, 25)
        ]
        self.assertIn("相控阵数字样机软件研制任务书", document_names)
        self.assertIn("相控阵数字样机软件开发计划", document_names)
        self.assertIn("相控阵数字样机软件设计说明", document_names)
        self.assertEqual("='ProgressTracking-ok'!H3+267", milestone["E25"].value)
        self.assertEqual("='ProgressTracking-ok'!M3+267", milestone["E27"].value)

    async def test_explicit_milestone_has_paired_report_number_and_no_meeting(self):
        report_types = await api_get_report_types(self.phase4_milestone_node_id, self.db)
        self.assertEqual(
            ["measure_milestone", "progress_milestone"],
            [item.id for item in report_types],
        )

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase4_id)
        ), data_only=False)
        measure = next(sheet for sheet in workbook if "测量与分析报告" in sheet.title and "里程碑" in sheet.title)
        progress = next(sheet for sheet in workbook if "进展报告" in sheet.title and "里程碑" in sheet.title)
        self.assertIn("测分报告-007", measure["E2"].value)
        self.assertIn("进展报告-007", progress["B2"].value)
        self.assertEqual(2, len([name for name in workbook.sheetnames if name.startswith("WorkReport")]))
        self.assertEqual(["LH008项目例会单", "LH009项目例会单"], [
            name for name in workbook.sheetnames if name.startswith("LH")
        ])

        next_phase_workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, self.project_id, self.phase5_id)
        ), data_only=False)
        self.assertIn("LH010项目例会单", next_phase_workbook.sheetnames)
        next_measure = next(
            sheet for sheet in next_phase_workbook
            if "测量与分析报告" in sheet.title
        )
        self.assertIn("测分报告-008", next_measure["E2"].value)


if __name__ == "__main__":
    unittest.main()
