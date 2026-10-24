import io
import tempfile
import unittest
from datetime import date
from pathlib import Path

import aiosqlite
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

import app.database as database
from app.services.cn_holidays import is_workday
from app.services.excel_builder.builder import build_excel
from app.services.importer import import_project
from app.services.workhour_engine import compute_node_metrics, working_days_in_range
from app.routers.projects import api_import


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_INPUT = ROOT / "项目相关信息表.xlsx"


def _sample_with_zero_workday_task() -> bytes:
    workbook = load_workbook(io.BytesIO(SAMPLE_INPUT.read_bytes()))
    schedule = workbook["项目进度计划"]
    target_row = next(
        row for row in range(2, schedule.max_row + 1)
        if str(schedule.cell(row, 1).value).strip() == "4.4.1"
    )
    schedule.cell(target_row, 3).value = date(2026, 4, 12)
    schedule.cell(target_row, 4).value = date(2026, 4, 12)
    invalid_file = io.BytesIO()
    workbook.save(invalid_file)
    return invalid_file.getvalue()


class WorkdayCalendarTests(unittest.TestCase):
    def test_2026_official_holidays_and_makeup_workdays(self):
        expected = {
            date(2026, 1, 3): False,
            date(2026, 1, 4): True,
            date(2026, 2, 14): True,
            date(2026, 2, 15): False,
            date(2026, 2, 16): False,
            date(2026, 2, 23): False,
            date(2026, 2, 24): True,
            date(2026, 2, 28): True,
            date(2026, 4, 12): False,
            date(2026, 5, 9): True,
            date(2026, 9, 20): True,
            date(2026, 9, 25): False,
            date(2026, 9, 27): False,
            date(2026, 10, 10): True,
        }
        self.assertEqual(expected, {day: is_workday(day) for day in expected})

    def test_february_24_one_day_task_is_fully_planned_and_completed(self):
        task_date = date(2026, 2, 24)

        self.assertEqual([task_date], working_days_in_range(task_date, task_date))
        metrics = compute_node_metrics(6.0, task_date, task_date, task_date)
        self.assertEqual(6.0, metrics.plan_should_complete)
        self.assertEqual(6.0, metrics.actual_complete)
        self.assertEqual(1.0, metrics.pct_complete)


class ImportWorkingDayValidationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_db_path = database.DB_PATH
        database.DB_PATH = str(Path(self.temp_dir.name) / "measurem.db")
        await database.init_db()
        self.db = await aiosqlite.connect(database.DB_PATH)

    async def asyncTearDown(self):
        await self.db.close()
        database.DB_PATH = self.old_db_path
        self.temp_dir.cleanup()

    async def test_zero_workday_reimport_is_rejected_before_existing_data_is_cleared(self):
        content = SAMPLE_INPUT.read_bytes()
        result = await import_project(self.db, content, SAMPLE_INPUT.name)

        with self.assertRaisesRegex(
            ValueError,
            r"没有有效工作日的任务：WBS 4\.4\.1 双周进度跟踪",
        ):
            await import_project(
                self.db, _sample_with_zero_workday_task(), SAMPLE_INPUT.name,
            )

        async with self.db.execute(
            "SELECT COUNT(*) FROM phases WHERE project_id=?",
            (result["project_id"],),
        ) as cursor:
            self.assertEqual(7, (await cursor.fetchone())[0])
        async with self.db.execute(
            """
            SELECT t.start_date, t.end_date
            FROM wbs_tasks t JOIN phases p ON p.id=t.phase_id
            WHERE p.project_id=? AND t.wbs='4.4.1'
            """,
            (result["project_id"],),
        ) as cursor:
            self.assertEqual(("2026-02-24", "2026-02-24"), await cursor.fetchone())

    async def test_import_api_returns_400_for_zero_workday_task(self):
        upload = UploadFile(
            file=io.BytesIO(_sample_with_zero_workday_task()),
            filename=SAMPLE_INPUT.name,
        )

        with self.assertRaises(HTTPException) as caught:
            await api_import(upload, self.db)

        self.assertEqual(400, caught.exception.status_code)
        self.assertIn("WBS 4.4.1 双周进度跟踪", caught.exception.detail)

    async def test_sample_phase4_metrics_and_workbook_use_february_24_as_workday(self):
        result = await import_project(
            self.db, SAMPLE_INPUT.read_bytes(), SAMPLE_INPUT.name,
        )
        async with self.db.execute(
            "SELECT id FROM phases WHERE project_id=? AND wbs='4'",
            (result["project_id"],),
        ) as cursor:
            phase_id = (await cursor.fetchone())[0]

        async with self.db.execute(
            """
            SELECT n.date,
                   ROUND(SUM(d.plan_should_complete), 2),
                   ROUND(SUM(d.actual_complete), 2)
            FROM task_node_data d
            JOIN wbs_tasks t ON t.id=d.task_id
            JOIN tracking_nodes n ON n.id=d.node_id
            WHERE t.phase_id=? AND n.node_type!='里程碑'
            GROUP BY n.id ORDER BY n.date
            """,
            (phase_id,),
        ) as cursor:
            self.assertEqual(
                [("2026-02-24", 61.0, 61.0), ("2026-03-11", 209.0, 209.0)],
                await cursor.fetchall(),
            )

        workbook = load_workbook(io.BytesIO(
            await build_excel(self.db, result["project_id"], phase_id)
        ), data_only=False)
        tracking = workbook["ProgressTracking-ok"]
        task_rows = {
            tracking.cell(row, 1).value: row
            for row in range(1, tracking.max_row + 1)
        }
        self.assertEqual(
            [1, 6, 1, 6],
            [
                tracking.cell(task_rows["4.4.1"], 9).value,
                tracking.cell(task_rows["4.4.2"], 9).value,
                tracking.cell(task_rows["4.4.1"], 13).value,
                tracking.cell(task_rows["4.4.2"], 13).value,
            ],
        )

        work_report = workbook["WorkReport20260224"]
        first_task_rows = {}
        for row in range(4, work_report.max_row + 1):
            wbs = work_report.cell(row, 2).value
            if wbs in {"4.4.1", "4.4.2"}:
                first_task_rows.setdefault(wbs, row)
        self.assertEqual({"4.4.1", "4.4.2"}, set(first_task_rows))
        for row in first_task_rows.values():
            self.assertEqual(1, work_report.cell(row, 9).value)
            self.assertGreater(
                sum(
                    float(work_report.cell(row, col).value or 0)
                    for col in range(12, work_report.max_column + 1)
                ),
                0,
            )


if __name__ == "__main__":
    unittest.main()
