"""
导入器：解析项目相关信息表.xlsx → 写入 DB，并批量计算
task_node_data + 生成 narrative 模板占位符。
"""
import io
import json
from datetime import datetime, date

import openpyxl
import aiosqlite

from app.services.workhour_engine import compute_plan_should, compute_actual_complete


# ── 工具函数 ─────────────────────────────────────────────────────────────────

def _safe(row, idx, default=None):
    try:
        v = row[idx]
        return v if v is not None else default
    except (IndexError, TypeError):
        return default


def _normalize_wbs(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        # 1.1 → "1.1"，去除浮点噪音
        return str(round(val, 6)).rstrip('0').rstrip('.')
    if isinstance(val, int):
        return str(val)
    s = str(val).strip()
    return s or None


def _to_date(val) -> date | None:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    return None


def _phase_wbs_of(wbs: str) -> str:
    return wbs.split('.')[0]


def _is_phase(wbs: str) -> bool:
    return '.' not in wbs


def _has_resources(resources_raw) -> bool:
    return bool(resources_raw and str(resources_raw).strip())


def _parse_resources(
    resources_raw, plan_hours: float, team_set: set[str] | None = None
) -> tuple[str, list[dict]]:
    """逗号分隔资源名 → JSON + list。仅保留团队成员，工时按内部人员均分。"""
    if not _has_resources(resources_raw):
        return "[]", []
    all_names = [n.strip() for n in str(resources_raw).split(",") if n.strip()]
    if not all_names or plan_hours <= 0:
        return "[]", []
    # 过滤非团队成员；未提供 team_set 时保留所有人
    names = [n for n in all_names if team_set is None or n in team_set]
    if not names:
        return "[]", []
    per_person = round(plan_hours / len(names), 2)
    resource_list = [{"name": n, "hours": per_person} for n in names]
    return json.dumps(resource_list, ensure_ascii=False), resource_list


# ── 解析 Excel ───────────────────────────────────────────────────────────────

def _parse_excel(content: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    # 项目信息（第2行）
    ws = wb["项目信息"]
    rows = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    r = rows[0] if rows else ()
    team_members_raw = str(_safe(r, 18) or "")
    project = {
        "name":         str(_safe(r, 0, "") or ""),
        "code":         str(_safe(r, 1, "") or ""),
        "manager":      str(_safe(r, 2) or ""),
        "cm_person":    str(_safe(r, 3) or ""),
        "qa_person":    str(_safe(r, 4) or ""),
        "team_members": team_members_raw,
    }
    # 团队成员集合，用于过滤非成员资源
    team_set: set[str] = {
        n.strip() for n in team_members_raw.split(",") if n.strip()
    }

    # 项目进度计划
    ws = wb["项目进度计划"]
    tasks = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _safe(row, 1)
        if not name:
            continue
        wbs = _normalize_wbs(_safe(row, 0))
        if not wbs:
            continue
        plan_hours = float(_safe(row, 6) or 0)
        tasks.append({
            "wbs":           wbs,
            "name":          str(name).strip(),
            "start_date":    _to_date(_safe(row, 2)),
            "end_date":      _to_date(_safe(row, 3)),
            "plan_hours":    plan_hours,
            "resources_raw": _safe(row, 7),
        })

    # 配置项清单（列：序号|名称|标识|版本|类型|实际页数|所属阶段WBS|评审完成日期|估算规模）
    config_items = []
    if "配置项清单" in wb.sheetnames:
        ws = wb["配置项清单"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = _safe(row, 1)
            if not name:
                continue
            seq       = _safe(row, 0)
            phase_raw = _safe(row, 6)
            actual_pg = _safe(row, 5)
            plan_pg   = _safe(row, 8)   # I列：估算规模
            review_dt = _safe(row, 7)   # H列：评审完成日期
            config_items.append({
                "seq":          int(seq) if seq is not None else None,
                "name":         str(name).strip(),
                "identifier":   str(_safe(row, 2) or ""),
                "version":      str(_safe(row, 3) or ""),
                "item_type":    str(_safe(row, 4) or ""),
                "actual_pages": int(actual_pg) if actual_pg is not None else None,
                "plan_pages":   int(plan_pg) if plan_pg is not None else None,
                "phase_wbs":    _normalize_wbs(phase_raw),
                "review_date":  _to_date(review_dt),
            })

    # 代码模块（列：模块序号|模块名称|实际行数|所属阶段WBS|完成日期|估算规模）
    code_modules = []
    if "代码模块" in wb.sheetnames:
        ws = wb["代码模块"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            loc = _safe(row, 2)
            if loc is None:
                continue
            seq           = _safe(row, 0)
            name          = _safe(row, 1)
            phase_raw     = _safe(row, 3)  # D列：所属阶段WBS
            completion_dt = _safe(row, 4)  # E列：完成日期
            plan_l        = _safe(row, 5)  # F列：估算规模
            code_modules.append({
                "seq":             int(seq) if seq is not None else None,
                "name":            str(name).strip() if name else "",
                "actual_loc":      int(loc),
                "plan_loc":        int(plan_l) if plan_l is not None else None,
                "phase_wbs":       _normalize_wbs(phase_raw),
                "completion_date": _to_date(completion_dt),
            })

    return {"project": project, "tasks": tasks, "team_set": team_set,
            "config_items": config_items, "code_modules": code_modules}


# ── 主入口 ───────────────────────────────────────────────────────────────────

async def import_project(
    db: aiosqlite.Connection,
    content: bytes,
    source_filename: str | None = None,
) -> dict:
    data     = _parse_excel(content)
    p        = data["project"]
    tasks    = data["tasks"]
    team_set = data["team_set"]

    # 1. Upsert 项目
    async with db.execute("SELECT id FROM projects WHERE code=?", (p["code"],)) as cur:
        existing = await cur.fetchone()

    if existing:
        project_id = existing[0]
        await _clear_project_data(db, project_id)
        await db.execute(
            "UPDATE projects SET name=?, manager=?, cm_person=?, qa_person=?, team_members=?, "
            "source_filename=?, created_at=datetime('now','localtime') WHERE id=?",
            (p["name"], p["manager"], p["cm_person"], p["qa_person"], p["team_members"],
             source_filename, project_id))
    else:
        async with db.execute(
            "INSERT INTO projects "
            "(name, code, manager, cm_person, qa_person, team_members, source_filename) "
            "VALUES (?,?,?,?,?,?,?)",
            (p["name"], p["code"], p["manager"], p["cm_person"], p["qa_person"],
             p["team_members"], source_filename)) as cur:
            project_id = cur.lastrowid

    # 2. 创建阶段（顶层 WBS 整数行）
    phase_rows = [t for t in tasks if _is_phase(t["wbs"])]
    phase_id_map: dict[str, int] = {}   # wbs_str → phase_id

    for pt in phase_rows:
        async with db.execute(
            "INSERT INTO phases (project_id, wbs, name, start_date, end_date) VALUES (?,?,?,?,?)",
            (project_id, pt["wbs"], pt["name"],
             pt["start_date"].isoformat() if pt["start_date"] else None,
             pt["end_date"].isoformat() if pt["end_date"] else None)) as cur:
            phase_id_map[pt["wbs"]] = cur.lastrowid

    # 3. 推断 tracking nodes（从任务名称含"进度跟踪"的行提取日期和类型）
    #    阶段1无此类任务，使用阶段 end_date 创建默认阶段节点
    node_id_map: dict[tuple[int, str], int] = {}       # (phase_id, date_iso) → node_id
    phase_node_map: dict[int, list[dict]] = {pid: [] for pid in phase_id_map.values()}

    for t in tasks:
        if "进度跟踪" not in t["name"]:
            continue
        pw = _phase_wbs_of(t["wbs"])
        phase_id = phase_id_map.get(pw)
        if phase_id is None:
            continue
        node_date = t["end_date"] or t["start_date"]
        if node_date is None:
            continue
        node_type = "双周" if "双周" in t["name"] else "阶段"
        date_iso  = node_date.isoformat()
        key = (phase_id, date_iso)
        if key not in node_id_map:
            async with db.execute(
                "INSERT INTO tracking_nodes (phase_id, date, node_type) VALUES (?,?,?)",
                (phase_id, date_iso, node_type)) as cur:
                nid = cur.lastrowid
            node_id_map[key] = nid
            phase_node_map[phase_id].append({"id": nid, "date": node_date, "type": node_type})

    # 为没有显式"进度跟踪"任务的阶段（如阶段1）创建默认阶段节点
    for wbs, phase_id in phase_id_map.items():
        if not phase_node_map[phase_id]:
            pt = next((t for t in phase_rows if t["wbs"] == wbs), None)
            if pt and pt["end_date"]:
                date_iso = pt["end_date"].isoformat()
                async with db.execute(
                    "INSERT INTO tracking_nodes (phase_id, date, node_type) VALUES (?,?,?)",
                    (phase_id, date_iso, "阶段")) as cur:
                    nid = cur.lastrowid
                node_id_map[(phase_id, date_iso)] = nid
                phase_node_map[phase_id].append({"id": nid, "date": pt["end_date"], "type": "阶段"})

    # 4. 插入 wbs_tasks（阶段行以外的所有行）
    leaf_by_phase: dict[int, list[dict]] = {pid: [] for pid in phase_id_map.values()}

    for t in tasks:
        if _is_phase(t["wbs"]):
            continue
        pw = _phase_wbs_of(t["wbs"])
        phase_id = phase_id_map.get(pw)
        if phase_id is None:
            continue

        resources_json, resource_list = _parse_resources(t["resources_raw"], t["plan_hours"], team_set)

        async with db.execute(
            "INSERT INTO wbs_tasks (phase_id, wbs, name, start_date, end_date, plan_hours, resources) "
            "VALUES (?,?,?,?,?,?,?)",
            (phase_id, t["wbs"], t["name"],
             t["start_date"].isoformat() if t["start_date"] else None,
             t["end_date"].isoformat() if t["end_date"] else None,
             t["plan_hours"], resources_json)) as cur:
            task_id = cur.lastrowid

        # 只有叶节点（有资源且有有效日期）才计算 task_node_data
        if resource_list and t["plan_hours"] > 0 and t["start_date"] and t["end_date"]:
            leaf_by_phase[phase_id].append({
                "id":         task_id,
                "wbs":        t["wbs"],
                "plan_hours": t["plan_hours"],
                "start_date": t["start_date"],
                "end_date":   t["end_date"],
                "resources":  resource_list,
                "is_phase1":  (pw == "1"),
            })

    # 5. 批量计算 task_node_data（三态逻辑：plan_should → actual → pct）
    for phase_id, leaf_tasks in leaf_by_phase.items():
        nodes = phase_node_map.get(phase_id, [])
        for task in leaf_tasks:
            ts, te = task["start_date"], task["end_date"]
            is_p1  = task["is_phase1"]

            for node_info in nodes:
                nd  = node_info["date"]
                nid = node_info["id"]

                plan_should = compute_plan_should(task["plan_hours"], ts, te, nd)
                actual, pct = compute_actual_complete(
                    task["plan_hours"], plan_should, ts, te, nd, is_phase1=is_p1,
                )

                await db.execute(
                    "INSERT OR REPLACE INTO task_node_data "
                    "(task_id, node_id, plan_should_complete, pct_complete, actual_complete) "
                    "VALUES (?,?,?,?,?)",
                    (task["id"], nid, plan_should, pct, actual))

    # 6. 生成 narrative 模板
    total_nodes = 0
    for phase_id, nodes in phase_node_map.items():
        pw        = next((w for w, pid in phase_id_map.items() if pid == phase_id), "")
        is_phase1 = (pw == "1")
        for node_index, node_info in enumerate(nodes):
            await _generate_narratives(
                db,
                node_info["id"],
                node_info["type"],
                is_phase1,
                phase_wbs=pw,
                is_first_node=(node_index == 0),
            )
            total_nodes += 1

    # 7. 写入配置项清单
    for item in data["config_items"]:
        plan_pg = item["plan_pages"] if item["plan_pages"] is not None else item["actual_pages"]
        await db.execute(
            "INSERT INTO config_items "
            "(project_id, seq, name, identifier, version, item_type, actual_pages, plan_pages, phase_wbs, review_date) "
            "VALUES (?,?,?,?,?,?,?,?,?,?)",
            (project_id, item["seq"], item["name"], item["identifier"],
             item["version"], item["item_type"], item["actual_pages"],
             plan_pg, item["phase_wbs"],
             item["review_date"].isoformat() if item["review_date"] else None))

    # 8. 写入代码模块
    for mod in data["code_modules"]:
        plan_l = mod["plan_loc"] if mod["plan_loc"] is not None else mod["actual_loc"]
        await db.execute(
            "INSERT INTO code_modules "
            "(project_id, seq, name, actual_loc, plan_loc, phase_wbs, completion_date) "
            "VALUES (?,?,?,?,?,?,?)",
            (project_id, mod["seq"], mod["name"], mod["actual_loc"],
             plan_l, mod["phase_wbs"],
             mod["completion_date"].isoformat() if mod["completion_date"] else None))

    await db.commit()

    async with db.execute(
        "SELECT COUNT(*) FROM wbs_tasks WHERE phase_id IN "
        "(SELECT id FROM phases WHERE project_id=?)", (project_id,)) as cur:
        tasks_count = (await cur.fetchone())[0]

    return {
        "project_id":    project_id,
        "project_name":  p["name"],
        "phases_count":  len(phase_id_map),
        "nodes_count":   total_nodes,
        "tasks_count":   tasks_count,
    }


# ── 清空已有数据 ──────────────────────────────────────────────────────────────

async def _clear_project_data(db: aiosqlite.Connection, project_id: int):
    await db.execute(
        "DELETE FROM narratives WHERE tracking_node_id IN "
        "(SELECT id FROM tracking_nodes WHERE phase_id IN "
        "(SELECT id FROM phases WHERE project_id=?))", (project_id,))
    await db.execute(
        "DELETE FROM task_node_data WHERE task_id IN "
        "(SELECT id FROM wbs_tasks WHERE phase_id IN "
        "(SELECT id FROM phases WHERE project_id=?))", (project_id,))
    await db.execute(
        "DELETE FROM task_node_data WHERE node_id IN "
        "(SELECT id FROM tracking_nodes WHERE phase_id IN "
        "(SELECT id FROM phases WHERE project_id=?))", (project_id,))
    await db.execute(
        "DELETE FROM wbs_tasks WHERE phase_id IN "
        "(SELECT id FROM phases WHERE project_id=?)", (project_id,))
    await db.execute(
        "DELETE FROM tracking_nodes WHERE phase_id IN "
        "(SELECT id FROM phases WHERE project_id=?)", (project_id,))
    await db.execute("DELETE FROM phases WHERE project_id=?", (project_id,))
    await db.execute("DELETE FROM config_items WHERE project_id=?", (project_id,))
    await db.execute("DELETE FROM code_modules WHERE project_id=?", (project_id,))


# ── Narrative 模板生成 ────────────────────────────────────────────────────────

_MEASURE_SECTIONS = [
    "deviation_tasks", "scale_analysis", "workload_analysis",
    "schedule_analysis", "req_analysis", "author",
]
_PROGRESS_PHASE_SECTIONS = [
    "change_summary", "progress_narrative", "deviation_stats",
    "overall", "risk_issues", "stakeholder", "resources_plan",
    "training", "issue_tracking", "current_summary", "next_plan", "author",
]
_PROGRESS_MILESTONE_SECTIONS = [
    "change_summary", "progress_narrative", "plan_completion",
    "risk_issues", "qa_work", "cm_work",
    "issue_tracking", "current_summary", "next_plan", "author",
]
_MEETING_SECTIONS = ["time", "location", "host", "attendees", "agenda"]


async def _generate_narratives(
    db: aiosqlite.Connection,
    node_id: int,
    node_type: str,
    is_phase1: bool,
    *,
    phase_wbs: str = "",
    is_first_node: bool = False,
):
    rows: list[tuple[str, str]] = []   # (report_type, section)

    if is_phase1:
        for s in ["deviation_tasks", "workload_analysis", "author"]:
            rows.append(("measure_phase1", s))
        for s in ["progress_narrative", "risk_issues", "current_summary", "next_plan", "author"]:
            rows.append(("progress_phase1", s))
        # 阶段1固定 3 个例会
        for prefix in ["lh001", "lh002", "lh003"]:
            for s in _MEETING_SECTIONS:
                rows.append(("meeting", f"{prefix}_{s}"))
    elif node_type == "双周":
        for s in _MEASURE_SECTIONS:
            rows.append(("measure_biweekly", s))
        for s in _PROGRESS_PHASE_SECTIONS:
            rows.append(("progress_biweekly", s))
        for s in _MEETING_SECTIONS:
            rows.append(("meeting", s))
    else:   # 阶段节点
        for s in _MEASURE_SECTIONS:
            rows.append(("measure_phase", s))
            rows.append(("measure_milestone", s))
        for s in _PROGRESS_PHASE_SECTIONS:
            rows.append(("progress_phase", s))
        for s in _PROGRESS_MILESTONE_SECTIONS:
            rows.append(("progress_milestone", s))
        for s in _MEETING_SECTIONS:
            rows.append(("meeting", s))

    if phase_wbs == "2" and is_first_node:
        for s in _MEETING_SECTIONS:
            rows.append(("meeting", f"planning_{s}"))

    for report_type, section in rows:
        await db.execute(
            "INSERT OR IGNORE INTO narratives (tracking_node_id, report_type, section, content) "
            "VALUES (?,?,?,?)",
            (node_id, report_type, section, ""))
