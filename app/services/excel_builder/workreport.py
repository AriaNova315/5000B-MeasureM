"""
WorkReport Sheet 构建器。
固定 11 列（A-K）+ 每个日历日一列（全阶段范围）。
按人员分组，每人先按 WBS 排序的任务列表。
工时记录只填写到 node_date 截止日，之后日期列留空。
"""
from collections import OrderedDict
from datetime import date, timedelta

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from .utils import (
    _w, _mc, _wmc, _s, border_range,
    F_WR_HDR, F_WR_BODY, A_CTR, A_CTR_W, A_LEFT, B_THIN,
    FILL_SECTION, FILL_GRID_HEADER, FILL_BODY, fill_range, set_outer_border,
    adaptive_row_height,
    set_col_width, set_row_height,
)

_N_FIXED = 11


def build_workreport(
    wb,
    sheet_name: str,
    tasks: list,        # [{id, wbs, name, start_date, end_date, plan_hours, resources}]
    phase: dict,        # {wbs, name, start_date, end_date}
    phase_nodes: list,  # [{id, date, type}]
    tnd: dict,          # {(task_id, node_id): {plan_should, pct, actual}}
    node_date: date,    # 本张 WorkReport 的工时截止日
    hour_dists: dict,   # {(task_id, person_name): {date: hours}} 预算分布
) -> dict[int, list[int]]:
    """
    生成一张 WorkReport。
    返回 task_wr_rows: {task_id: [行号列表]}，供 ProgressTracking 写引用公式。
    """
    ws = wb.create_sheet(sheet_name)

    # 日期列：全阶段范围（含节点截止日之后的列，但那些列工时留空）
    all_dates: list[date] = []
    d = phase["start_date"]
    while d <= phase["end_date"]:
        all_dates.append(d)
        d += timedelta(days=1)
    n_dates = len(all_dates)

    # 找到与 node_date 对应的当前节点 ID
    current_node = next((n for n in phase_nodes if n["date"] == node_date), None)
    current_nid  = current_node["id"] if current_node else None

    # ── 列宽 ─────────────────────────────────────────────────────────────────
    fixed_widths = [8, 8, 28, 11, 11, 20, 11, 11, 10, 10, 10]
    for i, w in enumerate(fixed_widths, 1):
        set_col_width(ws, i, w)
    for i in range(_N_FIXED + 1, _N_FIXED + n_dates + 1):
        set_col_width(ws, i, 4)

    # ── Row 1: 表头 ───────────────────────────────────────────────────────────
    headers = ["姓名", "WBS", "任务名称", "计划开始时间", "计划完成时间",
               "资源名称", "实际开始时间", "实际完成时间",
               "任务完成百分比", "实际已完成工作量", "个人实际总工时", "日工时记录"]
    row1 = headers + [None] * (n_dates - 1)
    ws.append(row1)
    for c in range(1, _N_FIXED + 2):
        _s(ws, 1, c, font=F_WR_HDR, alignment=A_CTR_W)
    if n_dates > 1:
        _mc(ws, 1, _N_FIXED + 1, 1, _N_FIXED + n_dates)
    set_row_height(ws, 1, 30)

    # ── Row 2: 星期几 ─────────────────────────────────────────────────────────
    _WEEKDAY_CN = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]
    row2 = [None] * _N_FIXED + [_WEEKDAY_CN[d.weekday()] for d in all_dates]
    ws.append(row2)
    for c in range(_N_FIXED + 1, _N_FIXED + n_dates + 1):
        _s(ws, 2, c, font=F_WR_HDR, alignment=A_CTR)

    # ── Row 3: 日期值（YYYY/M/D，无前导零）────────────────────────────────────
    row3 = [None] * _N_FIXED + list(all_dates)
    ws.append(row3)
    for c in range(_N_FIXED + 1, _N_FIXED + n_dates + 1):
        _s(ws, 3, c, font=F_WR_HDR, alignment=A_CTR, number_format="YYYY/M/D")
    # 固定字段表头纵向占满三行；日期区域继续保留“总表头 / 星期 / 日期”三级结构。
    for c in range(1, _N_FIXED + 1):
        _mc(ws, 1, c, 3, c)
    set_row_height(ws, 2, 24)
    set_row_height(ws, 3, 24)

    last_date_col = get_column_letter(_N_FIXED + n_dates)

    # ── 收集 (person, task) ───────────────────────────────────────────────────
    person_order: list[str] = []
    person_tasks: dict[str, list] = OrderedDict()

    for task in tasks:
        if not task.get("resources") or not task.get("start_date") or not task.get("end_date"):
            continue
        for res in task["resources"]:
            person = res["name"]
            dist   = hour_dists.get((task["id"], person), {})

            pct = 0.0
            task_actual = 0.0
            if current_nid and (task["id"], current_nid) in tnd:
                td          = tnd[(task["id"], current_nid)]
                pct         = td["pct"]
                task_actual = td["actual"]
            resources_str = ",".join(r["name"] for r in task["resources"])

            entry = {
                "task_id":       task["id"],
                "wbs":           task["wbs"],
                "name":          task["name"],
                "plan_start":    task["start_date"],
                "plan_end":      task["end_date"],
                "plan_hours":    task["plan_hours"],
                "resources_str": resources_str,
                "pct":           pct,
                "actual_comp":   round(task_actual, 2),
                "dist":          dist,
            }

            if person not in person_tasks:
                person_order.append(person)
                person_tasks[person] = []
            person_tasks[person].append(entry)

    # ── 写入数据行 ────────────────────────────────────────────────────────────
    task_wr_rows: dict[int, list[int]] = {}  # task_id → [行号]
    task_first_wr_row: dict[int, int] = {}
    dr = 4
    for person in person_order:
        entries = person_tasks[person]
        for i, e in enumerate(entries):
            show_name = person if i == 0 else None
            first_task_row = task_first_wr_row.get(e["task_id"])
            pct_value = e["pct"] if first_task_row is None else f"=I{first_task_row}"
            row = [
                show_name,
                e["wbs"],
                e["name"],
                e["plan_start"],
                e["plan_end"],
                e["resources_str"],
                None,           # G: 实际开始时间（公式）
                None,           # H: 实际完成时间（公式）
                pct_value,
                f'={e["plan_hours"]}*I{dr}',
                None,           # K: 个人实际总工时（公式）
            ]
            # 日工时列：截止 node_date，之后留空
            for d in all_dates:
                h = e["dist"].get(d) if d <= node_date else None
                row.append(h if h else None)
            ws.append(row)

            # G列：数组公式，找第一个有工时的日期
            g_formula = (
                f"=IF(ISNUMBER(INDEX($L$3:${last_date_col}$3,"
                f"MATCH(TRUE,L{dr}:{last_date_col}{dr}>0,))),"
                f"INDEX($L$3:${last_date_col}$3,"
                f"MATCH(TRUE,L{dr}:{last_date_col}{dr}>0,)),\"\")"
            )
            ws.cell(dr, 7).value = ArrayFormula(f"G{dr}", g_formula)

            # H列：找最后一个有工时的日期（仅完成100%时填值）
            ws.cell(dr, 8).value = (
                f"=IF(AND(ISNUMBER(LOOKUP(1,0/(L{dr}:{last_date_col}{dr}<>\"\"),"
                f"$L$3:${last_date_col}$3)),I{dr}=1),"
                f"LOOKUP(1,0/(L{dr}:{last_date_col}{dr}<>\"\"),"
                f"$L$3:${last_date_col}$3),\"\")"
            )

            # K列：个人实际总工时
            ws.cell(dr, 11).value = f"=SUM(L{dr}:{last_date_col}{dr})"

            # 样式
            for c in range(1, _N_FIXED + 1):
                _s(ws, dr, c, font=F_WR_BODY, alignment=A_LEFT)
            for dc in (4, 5, 7, 8):
                _s(ws, dr, dc, number_format="YYYY/MM/DD")
            _s(ws, dr, 9, number_format="0.00%")
            for c in range(_N_FIXED + 1, _N_FIXED + n_dates + 1):
                _s(ws, dr, c, font=F_WR_BODY, alignment=A_CTR)

            set_row_height(
                ws, dr,
                adaptive_row_height(
                    f"{e['name']}\n{e['resources_str']}",
                    min_height=24, chars_per_line=36, max_height=60,
                ),
            )

            # 记录行号（供 ProgressTracking 引用）
            task_wr_rows.setdefault(e["task_id"], []).append(dr)
            task_first_wr_row.setdefault(e["task_id"], dr)
            dr += 1

    border_range(ws, 1, 1, dr - 1, _N_FIXED + n_dates)
    fill_range(ws, 1, 1, dr - 1, _N_FIXED + n_dates, FILL_BODY)
    fill_range(ws, 1, 1, 1, _N_FIXED + n_dates, FILL_SECTION)
    fill_range(ws, 1, 1, 3, _N_FIXED, FILL_SECTION)
    fill_range(ws, 2, _N_FIXED + 1, 3, _N_FIXED + n_dates, FILL_GRID_HEADER)
    set_outer_border(ws, 1, 1, dr - 1, _N_FIXED + n_dates)
    return task_wr_rows
