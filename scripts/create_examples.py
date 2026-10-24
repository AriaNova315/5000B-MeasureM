"""Create distributable input workbooks from synthetic data, never a real project."""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = {
    "项目信息": ["项目名称", "项目代号", "项目组长", "配置管理员", "质量管理员", "估计轮次",
             "文档估计人1", "文档估计人2", "文档估计人3", "文档估计人4", "代码估计人1",
             "代码估计人2", "代码估计人3", "代码估计人4", "CM估计人", "QA估计人", "管理估计人", "评审估计人", "项目组成员\n(逗号分隔)"],
    "项目进度计划": ["WBS", "任务名称", "计划开始时间", "计划完成时间", "关联配置项", "工期(天)", "工作量(人时)", "资源名称\n(逗号分隔)"],
    "里程碑": ["序号", "里程碑名称", "里程碑日期", "累计截止阶段WBS", "累计截止阶段名称", "生成测量分析报告", "生成进展报告"],
    "配置项清单": ["序号", "配置项名称", "配置项标识", "版本", "类型\n(文档/代码/其他)", "实际页数\n(仅文档类填写)", "所属阶段WBS", "评审完成日期", "估算规模"],
    "基线组成": ["基线名称", "基线标识", "基线版本", "包含配置项名称", "配置项标识", "配置项版本"],
    "代码模块": ["模块序号", "模块名称", "实际行数(LOC)", "所属阶段WBS", "完成日期", "估算规模"],
}


def create_workbook(demo=False):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in HEADERS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 36
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
            cell.alignment = Alignment(wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = 22
    if demo:
        wb["项目信息"].append(["离线验收演示项目", "DEMO", "演示甲", "演示乙", "演示丙"] + [None] * 13 + ["演示甲,演示乙,演示丙"])
        ws = wb["项目进度计划"]
        def task(wbs, name, start, end, hours=0, resources=None):
            ws.append([wbs, name, date(2026, 3, start), date(2026, 3, end), None, None, hours, resources])
        task("1", "系统分析与设计阶段", 2, 6)
        task("1.1", "软件工程活动", 2, 6)
        task("1.1.1", "编写演示任务书", 2, 6, 20, "演示甲,演示乙")
        task("1.2", "阶段例会", 6, 6, 2, "演示甲,演示乙,演示丙")
        task("2", "需求分析阶段", 9, 20)
        task("2.2", "软件工程活动", 9, 20)
        task("2.2.1", "编写演示需求", 9, 20, 40, "演示甲,演示乙")
        task("2.3", "评审活动", 19, 20)
        task("2.3.1", "评审演示需求", 19, 20, 6, "演示丙")
        task("2.4.1", "双周进度跟踪", 13, 13, 1, "演示甲")
        task("2.4.2", "阶段进度跟踪", 20, 20, 1, "演示甲")
        task("2.5.1", "双周例会", 13, 13, 2, "演示甲,演示乙")
        task("2.5.2", "阶段例会", 20, 20, 2, "演示甲,演示乙,演示丙")
        wb["里程碑"].append([1, "演示需求里程碑", date(2026, 3, 20), "2", "需求分析阶段", "是", "是"])
        wb["配置项清单"].append([1, "演示任务书", "DEMO-DOC-01", "1.0", "文档", 8, "1", date(2026, 3, 6), 8])
        wb["配置项清单"].append([2, "演示需求说明书", "DEMO-DOC-02", "1.0", "文档", 16, "2", date(2026, 3, 20), 15])
        wb["代码模块"].append([1, "演示模块", 100, "2", date(2026, 3, 20), 100])
    for ws in wb:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
    return wb


if __name__ == "__main__":
    destination = Path(__file__).resolve().parents[1] / "examples"
    destination.mkdir(exist_ok=True)
    for name, demo in [("input-template.xlsx", False), ("demo-project.xlsx", True)]:
        target = destination / name
        if target.exists():
            raise SystemExit(f"Refusing to overwrite {target}")
        wb = create_workbook(demo)
        wb.save(target)
        wb.close()
        print(target)
