"""Offline core-flow verification using a temporary DB and synthetic input only."""
import asyncio
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

# Source ZIPs include a manifest; a development checkout need not have one.
manifest_path = ROOT / "MANIFEST.json"
if manifest_path.is_file():
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for name, expected in manifest.items():
        target = (ROOT / name).resolve()
        if not target.is_relative_to(ROOT) or hashlib.sha256(target.read_bytes()).hexdigest() != expected:
            raise SystemExit(f"Delivery file checksum mismatch: {name}")
    print(f"PASS: {len(manifest)} delivery file checksums")

import aiosqlite
from openpyxl import load_workbook

import app.database as database
from app.models import NarrativeUpdate
from app.routers import export
from app.routers.projects import (
    api_get_narratives, api_get_phases, api_list_projects, api_update_narrative,
)
from app.services.importer import import_project


async def verify():
    old_db, old_exports = database.DB_PATH, export.GENERATED_DIR
    try:
        with tempfile.TemporaryDirectory(prefix="measurem-verify-") as temporary:
            database.DB_PATH = str(Path(temporary) / "verify.db")
            export.GENERATED_DIR = str(Path(temporary) / "generated")
            await database.init_db()
            await database.init_db()
            async with aiosqlite.connect(database.DB_PATH) as db:
                tables = await (await db.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
                )).fetchall()
                assert len(tables) == 8
                assert await api_list_projects(db) == []
                example = ROOT / "examples" / "demo-project.xlsx"
                result = await import_project(db, example.read_bytes(), example.name)
                project_id = result["project_id"]
                assert len(await api_list_projects(db)) == 1
                phases = await api_get_phases(project_id, db)
                assert len(phases) == 2
                assert any(n.node_type == "里程碑" for ph in phases for n in ph.nodes)
                node_id = phases[0].nodes[0].id
                marker = "离线验收叙述覆盖成功"
                await api_update_narrative(node_id, "progress_phase1", "current_summary", NarrativeUpdate(content=marker), db)
                sections = await api_get_narratives(node_id, "progress_phase1", db)
                assert any(s.content == marker and s.is_overridden for s in sections)
                for phase in phases:
                    response = await export.api_export_phase(project_id, phase.id, db)
                    wb = load_workbook(io.BytesIO(response.body), data_only=False)
                    try:
                        assert "ProgressTracking-ok" in wb.sheetnames
                        assert any(n.startswith("WorkReport") for n in wb.sheetnames)
                        assert wb.calculation.fullCalcOnLoad
                        if phase.wbs == "1":
                            assert wb["ProgressTracking-ok"].max_column == 5
                            assert any(c.value == marker for ws in wb for row in ws for c in row)
                        else:
                            assert any("里程碑" in n for n in wb.sheetnames)
                        print(f"PASS phase {phase.wbs}: {len(wb.sheetnames)} sheets")
                    finally:
                        wb.close()
                await api_update_narrative(node_id, "progress_phase1", "current_summary", NarrativeUpdate(content=""), db)
                sections = await api_get_narratives(node_id, "progress_phase1", db)
                assert all(not s.is_overridden for s in sections if s.section == "current_summary")
            print("PASS: empty DB, import, list, narrative edit/reset and all demo exports.")
            print("Real application data was not opened. Excel recalculation requires desktop validation.")
    finally:
        database.DB_PATH, export.GENERATED_DIR = old_db, old_exports


if __name__ == "__main__":
    asyncio.run(verify())
